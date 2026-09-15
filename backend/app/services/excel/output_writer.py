from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.excel_models import OutputReport, OutputValues
from app.services.excel.graduation_parser import TARGET_MOS, TARGET_TITLE, _find_headers
from app.services.imports.validation import build_header_index


def generate_output(source_path: str, result_map: dict[tuple[str, int], OutputValues], destination_path: str) -> OutputReport:
    workbook = load_workbook(Path(source_path), data_only=False)
    changes: list[dict[str, object]] = []
    for worksheet in workbook.worksheets:
        _, headers = _find_headers(worksheet)
        index = build_header_index(headers)
        if TARGET_MOS.casefold() not in index or TARGET_TITLE.casefold() not in index:
            continue
        mos_column = index[TARGET_MOS.casefold()] + 1
        title_column = index[TARGET_TITLE.casefold()] + 1
        for (sheet_name, source_row), values in result_map.items():
            if sheet_name != worksheet.title:
                continue
            worksheet.cell(source_row, mos_column).value = values.mos_value
            worksheet.cell(source_row, title_column).value = values.title_value
            changes.append({"sheet": sheet_name, "row": source_row, "mosValue": values.mos_value, "titleValue": values.title_value})
    output = Path(destination_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return OutputReport(str(output), tuple(changes))
