from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.excel_models import GraduationRow, GraduationWorkbookReport
from app.domain.matching.normalization import normalize_nim
from app.services.imports.validation import build_header_index

TARGET_MOS = "Sertifikasi MOS (Diisi ITCC)"
TARGET_TITLE = "Title Microsoft (ITCC)"


def parse_graduation_workbook(path: str) -> GraduationWorkbookReport:
    workbook = load_workbook(Path(path), read_only=True, data_only=False)
    rows: list[GraduationRow] = []
    errors: list[str] = []
    for worksheet in workbook.worksheets:
        header_row_number, headers = _find_headers(worksheet)
        index = build_header_index(headers)
        if "nim" not in index or "nama" not in index:
            errors.append(f"Sheet {worksheet.title} tidak memiliki header NIM dan Nama.")
            continue
        if TARGET_MOS.casefold() not in index or TARGET_TITLE.casefold() not in index:
            errors.append(f"Sheet {worksheet.title} tidak memiliki dua kolom ITCC.")
            continue
        for source_row, values in enumerate(worksheet.iter_rows(min_row=header_row_number + 1, values_only=True), start=header_row_number + 1):
            nim = normalize_nim(values[index["nim"]] if index["nim"] < len(values) else None)
            nama = str(values[index["nama"]] or "").strip() if index["nama"] < len(values) else ""
            if not nim and not nama:
                continue
            program_code = ""
            for alias in ("kode prodi", "kode program studi", "program code"):
                if alias in index and index[alias] < len(values):
                    program_code = str(values[index[alias]] or "").strip()
                    break
            row_values = {headers[position]: values[position] if position < len(values) else None for position in range(len(headers))}
            rows.append(GraduationRow(worksheet.title, source_row, nim, nama, program_code, row_values))
    return GraduationWorkbookReport(tuple(workbook.sheetnames), tuple(rows), tuple(errors))


def _find_headers(worksheet: object) -> tuple[int, tuple[str, ...]]:
    for row_number, values in enumerate(worksheet.iter_rows(min_row=1, max_row=min(20, worksheet.max_row), values_only=True), start=1):
        index = build_header_index(values)
        if "nim" in index and "nama" in index:
            return row_number, tuple(str(value or "") for value in values)
    return 1, ()
