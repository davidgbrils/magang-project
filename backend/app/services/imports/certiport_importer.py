from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.imports.models import ImportIssue, ImportedRecord, ImportReport
from app.domain.matching.normalization import normalize_name, normalize_nim
from .validation import build_header_index, missing_headers

REQUIRED_HEADERS = ("Student / Employee ID", "First Name", "Last Name", "Exam", "Program Name", "Exam Date", "Result")
ACCEPTED_RESULTS = frozenset({"pass", "fail", "inco"})


def parse_certiport(file_path: str) -> ImportReport:
    workbook = load_workbook(Path(file_path), read_only=True, data_only=True)
    worksheet = workbook.active
    header_row = next(worksheet.iter_rows(min_row=4, max_row=4, values_only=True), ())
    missing = missing_headers(header_row, REQUIRED_HEADERS)
    headers = tuple(str(value or "") for value in header_row)
    if missing:
        return ImportReport("INVALID", headers, 0, (), tuple(ImportIssue(4, header, "MISSING_HEADER", f"Kolom {header} wajib ada.") for header in missing))
    index = build_header_index(header_row)
    records: list[ImportedRecord] = []
    errors: list[ImportIssue] = []
    warnings: list[ImportIssue] = []
    for row_number, values in enumerate(worksheet.iter_rows(min_row=5, values_only=True), start=5):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        raw_row = {headers[position]: values[position] if position < len(values) else None for position in range(len(headers))}
        result = str(values[index["result"]] or "").strip()
        if result.casefold() not in ACCEPTED_RESULTS:
            errors.append(ImportIssue(row_number, "Result", "INVALID_RESULT", "Result harus Pass, Fail, atau Inco."))
            continue
        first_name = raw_row.get("First Name") or ""
        last_name = raw_row.get("Last Name") or ""
        full_name = f"{first_name} {last_name}".strip()
        student_id = normalize_nim(values[index["student / employee id"]] if index["student / employee id"] < len(values) else None)
        if not student_id:
            warnings.append(ImportIssue(row_number, "Student / Employee ID", "MISSING_STUDENT_ID", "Record tetap dapat direview berdasarkan nama."))
        records.append(ImportedRecord(row_number, {"student_employee_id": student_id, "first_name": first_name, "last_name": last_name, "full_name": full_name, "full_name_normalized": normalize_name(full_name), "exam": raw_row.get("Exam"), "program_name": raw_row.get("Program Name"), "exam_date": raw_row.get("Exam Date"), "result": result}, raw_row))
    return ImportReport("VALID" if not errors else "INVALID", headers, len(records), tuple(records), tuple(errors), tuple(warnings))
