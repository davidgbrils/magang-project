from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.domain.imports.models import ImportIssue, ImportedRecord, ImportReport
from app.domain.matching.normalization import normalize_nim
from .validation import build_header_index, missing_headers

REQUIRED_HEADERS = ("NIM", "Nama", "Status")


def parse_sitasi(file_path: str) -> ImportReport:
    workbook = load_workbook(Path(file_path), read_only=True, data_only=True)
    worksheet = workbook.active
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    missing = missing_headers(header_row, REQUIRED_HEADERS)
    if missing:
        return ImportReport("INVALID", tuple(str(value or "") for value in header_row), 0, (), tuple(ImportIssue(1, header, "MISSING_HEADER", f"Kolom {header} wajib ada.") for header in missing))

    header_index = build_header_index(header_row)
    headers = tuple(str(value or "") for value in header_row)
    records: list[ImportedRecord] = []
    errors: list[ImportIssue] = []
    nims: Counter[str] = Counter()
    for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        raw_row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        nim = normalize_nim(values[header_index["nim"]] if header_index["nim"] < len(values) else None)
        if not nim:
            errors.append(ImportIssue(row_number, "NIM", "INVALID_NIM", "NIM wajib diisi."))
            continue
        nims[nim] += 1
        records.append(ImportedRecord(row_number, {"nim": nim, "nama": raw_row.get("Nama"), "status": raw_row.get("Status")}, raw_row))
    return ImportReport("VALID" if not errors else "INVALID", headers, len(records), tuple(records), tuple(errors), duplicate_count=sum(count - 1 for count in nims.values() if count > 1))
