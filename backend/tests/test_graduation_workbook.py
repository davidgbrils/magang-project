from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.domain.excel_models import OutputValues
from app.services.excel.graduation_parser import parse_graduation_workbook
from app.services.excel.output_writer import generate_output


def _create_graduation_workbook(path: Path) -> list[str]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    names = [f"S{i:02d}" for i in range(1, 15)]
    for name in names:
        sheet = workbook.create_sheet(name)
        sheet.append(["NIM", "Nama", "Kode Prodi", "Nilai", "Sertifikasi MOS (Diisi ITCC)", "Title Microsoft (ITCC)"])
        sheet.append(["202431001", "Ayu Putri", "31", "=1+1", None, None])
    workbook.save(path)
    return names


def test_output_preserves_fourteen_sheets_and_only_two_itcc_columns(tmp_path: Path) -> None:
    source = tmp_path / "graduation.xlsx"
    destination = tmp_path / "generated.xlsx"
    sheet_names = _create_graduation_workbook(source)

    parsed = parse_graduation_workbook(str(source))
    report = generate_output(str(source), {("S01", 2): OutputValues("MOS (Word 2019 - Lulus)", "MOS")}, str(destination))

    original = load_workbook(source, data_only=False)
    generated = load_workbook(destination, data_only=False)
    assert parsed.sheet_names == tuple(sheet_names)
    assert len(parsed.rows) == 14
    assert report.changes[0]["sheet"] == "S01"
    assert generated.sheetnames == sheet_names
    for sheet_name in sheet_names:
        before = original[sheet_name]
        after = generated[sheet_name]
        for row in range(1, before.max_row + 1):
            for column in range(1, before.max_column + 1):
                if sheet_name == "S01" and row == 2 and column in {5, 6}:
                    continue
                assert after.cell(row, column).value == before.cell(row, column).value
    assert generated["S01"].cell(2, 5).value == "MOS (Word 2019 - Lulus)"
    assert generated["S01"].cell(2, 6).value == "MOS"
