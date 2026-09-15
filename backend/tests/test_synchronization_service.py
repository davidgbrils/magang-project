from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services.synchronization_service import run_synchronization


def test_pipeline_generates_manifest_and_workbook_after_deterministic_match(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    destination = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S01"
    sheet.append(["NIM", "Nama", "Kode Prodi", "Sertifikasi MOS (Diisi ITCC)", "Title Microsoft (ITCC)"])
    sheet.append(["202431001", "Ayu Putri", "31", None, None])
    workbook.save(source)

    result = run_synchronization(
        str(source),
        [{"nim": "202431001", "status": "LULUS"}],
        [{"student_employee_id": "202431001", "full_name": "Ayu Putri", "program_name": "Microsoft Office Specialist", "exam": "Word 2019", "exam_date": "2026-01-01", "result": "Pass"}],
        str(destination),
    )

    generated = load_workbook(destination)
    assert result.status == "COMPLETED"
    assert result.manifest[0]["decision"] == "READY"
    assert generated["S01"].cell(2, 4).value == "MOS (Word 2019 - Lulus)"
    assert generated["S01"].cell(2, 5).value == "MOS"


def test_pipeline_does_not_generate_output_when_name_only_review_remains(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["NIM", "Nama", "Kode Prodi", "Sertifikasi MOS (Diisi ITCC)", "Title Microsoft (ITCC)"])
    sheet.append(["202431001", "Ayu Putri", "31", None, None])
    workbook.save(source)

    result = run_synchronization(
        str(source),
        [{"nim": "202431001", "status": "LULUS"}],
        [{"student_employee_id": "", "full_name": "Ayu Putri", "program_name": "Microsoft Office Specialist", "exam": "Word 2019", "exam_date": "2026-01-01", "result": "Pass"}],
    )

    assert result.status == "READY_FOR_REVIEW"
    assert result.output_path is None
